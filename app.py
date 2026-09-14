import io
import os
import re
import time
from flask import Flask, jsonify, request, send_from_directory, send_file
import resend
import fitz
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__)

DEMO_ITEMS = [
    (1,"3348901601201","Demo Luxury Item 01",24,112.50,2700.00),(2,"3348901601218","Demo Luxury Item 02",18,89.95,1619.10),
    (3,"3348901601225","Demo Luxury Item 03",36,145.00,5220.00),(4,"3348901601232","Demo Luxury Item 04",12,76.40,916.80),
    (5,"3348901601249","Demo Luxury Item 05",30,132.75,3982.50),(6,"3348901601256","Demo Luxury Item 06",20,98.20,1964.00),
    (7,"3348901601263","Demo Luxury Item 07",16,210.00,3360.00),(8,"3348901601270","Demo Luxury Item 08",40,64.50,2580.00),
    (9,"3348901601287","Demo Luxury Item 09",10,255.80,2558.00),(10,"3348901601294","Demo Luxury Item 10",28,118.30,3312.40),
    (11,"3348901601300","Demo Luxury Item 11",22,92.60,2037.20),(12,"3348901601317","Demo Luxury Item 12",14,174.25,2439.50),
    (13,"3348901601324","Demo Luxury Item 13",32,83.90,2684.80),(14,"3348901601331","Demo Luxury Item 14",26,149.50,3887.00),
    (15,"3348901601348","Demo Luxury Item 15",8,320.00,2560.00),(16,"3348901601355","Demo Luxury Item 16",44,57.75,2541.00),
    (17,"3348901601362","Demo Luxury Item 17",15,188.40,2826.00),(18,"3348901601379","Demo Luxury Item 18",19,104.90,1993.10),
    (19,"3348901601386","Demo Luxury Item 19",27,136.20,3677.40),(20,"3348901601393","Demo Luxury Item 20",11,225.00,2475.00)
]

@app.get("/")
def index():
    return send_from_directory(BASE_DIR, "nova.html")

@app.get("/nova.html")
def nova_page():
    return send_from_directory(BASE_DIR, "nova.html")

@app.get("/demo-files/<path:filename>")
def demo_files(filename):
    allowed = {"PO_3987563_DEMO.pdf", "INVOICE_INV-DEMO-260914.pdf", "ASN_RPA_TEMPLATE_DEMO.xlsx", "PO_7426158_DEMO.pdf", "INVOICE_INV-DEMO-50-260914.pdf", "ASN_RPA_DEMO_20260914_SHP-29718.xlsx", "PO_5843197_DEMO.pdf", "INVOICE_INV-DEMO-ERR-260914.pdf"}
    if filename not in allowed:
        return jsonify({"ok": False, "error": "File not found"}), 404
    return send_from_directory(BASE_DIR, filename)


@app.get("/ASN_RPA_DEMO_20260914_SHP-29641.xlsx")
def completed_rpa_demo():
    return send_from_directory(BASE_DIR, "ASN_RPA_DEMO_20260914_SHP-29641.xlsx", as_attachment=True)

@app.get("/health")
def health():
    return jsonify({"ok": True, "configured": bool(os.getenv("RESEND_API_KEY", "").strip())})

def pdf_text(file_storage):
    data = file_storage.read()
    file_storage.stream.seek(0)
    doc = fitz.open(stream=data, filetype="pdf")
    return "\n".join(page.get_text("text") for page in doc)

@app.post("/analyze-po")
def analyze_po():
    try:
        po_number = (request.form.get("po_number") or "").strip()
        if not re.fullmatch(r"\d{7}", po_number):
            return jsonify({"ok": False, "error": "Enter a valid 7-digit PO number."}), 400
        if po_number != "3987563":
            return jsonify({"ok": False, "error": f"PO {po_number} was not found in this controlled demo. Try 3987563."}), 404
        return jsonify({
            "ok": True,
            "po_number": "3987563",
            "currency": "EUR",
            "total_qty": 452,
            "total_value": 55333.80,
            "items": [
                {"line": r[0], "barcode": r[1], "description": r[2], "qty": r[3], "unit_price": r[4], "total": r[5]}
                for r in DEMO_ITEMS
            ]
        })
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

@app.post("/process-rpa")
def process_rpa():
    started = time.perf_counter()
    try:
        po_number = (request.form.get("po_number") or "").strip()
        invoice = request.files.get("invoice")
        if not re.fullmatch(r"\d{7}", po_number):
            return jsonify({"ok": False, "error": "A valid 7-digit PO number is required."}), 400
        if po_number != "3987563":
            return jsonify({"ok": False, "error": f"PO {po_number} was not found in this controlled demo. Try 3987563."}), 404
        if not invoice:
            return jsonify({"ok": False, "error": "Invoice file is required."}), 400

        inv_text = pdf_text(invoice)
        if "INV-DEMO-260914" not in inv_text or po_number not in inv_text:
            return jsonify({"ok": False, "error": "The uploaded invoice does not match the entered PO or demo invoice reference."}), 400

        # Controlled POC validation: the PO data is retrieved by PO number; every PO barcode must be present in the invoice.
        missing_inv = [r[1] for r in DEMO_ITEMS if r[1] not in inv_text]
        if missing_inv:
            return jsonify({"ok": False, "error": "One or more PO barcode lines could not be matched against the uploaded invoice."}), 400

        template = os.path.join(BASE_DIR, "ASN_RPA_TEMPLATE_DEMO.xlsx")
        wb = load_workbook(template)
        ws = wb[wb.sheetnames[0]]

        # Copy the template row style/structure down to 20 item rows.
        from copy import copy
        for row in range(3, 22):
            for col in range(1, 38):
                src = ws.cell(2, col)
                dst = ws.cell(row, col)
                if src.has_style:
                    dst._style = copy(src._style)
                if src.number_format:
                    dst.number_format = src.number_format
                if src.font:
                    dst.font = copy(src.font)
                if src.fill:
                    dst.fill = copy(src.fill)
                if src.border:
                    dst.border = copy(src.border)
                if src.alignment:
                    dst.alignment = copy(src.alignment)

        # A-AK fields. Barcode used; SKU intentionally blank. UDF4 cartons; UDF5 pallets blank.
        for idx, r in enumerate(DEMO_ITEMS, start=2):
            _, barcode, _, qty, unit_price, total = r
            values = {
                1:"Sequence", 2:po_number, 3:"INV-DEMO-260914", 4:"14/09/2026", 5:barcode,
                6:None, 7:None, 8:qty, 9:unit_price, 10:"EUR", 11:total, 14:"SHP-29641",
                25:"DAP", 26:"FRANCE", 28:186.4, 29:"UNITED ARAB EMIRATES", 30:"AIR",
                32:12, 33:None, 34:"MACWLL-PharaohsTeam@chalhoub.com", 35:"MACWLL-PharaohsTeam@chalhoub.com"
            }
            for col, val in values.items():
                ws.cell(idx, col).value = val

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        elapsed_ms = (time.perf_counter() - started) * 1000
        resp = send_file(out, as_attachment=True, download_name="ASN_RPA_DEMO_20260914_SHP-29641.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp.headers["X-Nova-Processing-MS"] = f"{elapsed_ms:.2f}"
        return resp
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.after_request
def nova_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    return response

@app.post("/send-email")
def send_email():
    try:
        api_key = os.getenv("RESEND_API_KEY", "").strip()
        if not api_key:
            return jsonify({"ok": False, "error": "RESEND_API_KEY is not configured on the server."}), 500

        resend.api_key = api_key
        payload = request.get_json(force=True) or {}
        recipient = (payload.get("to") or "").strip()
        subject = (payload.get("subject") or "Nova Shipment Update").strip()
        html = payload.get("html") or "<p>Nova shipment update.</p>"

        if "@" not in recipient:
            return jsonify({"ok": False, "error": "Invalid recipient email address."}), 400

        sender_name = os.getenv("NOVA_SENDER_NAME", "Nova | Digital Shipping Agent").strip()
        sender_email = os.getenv("NOVA_SENDER_EMAIL", "onboarding@resend.dev").strip()

        result = resend.Emails.send({
            "from": f"{sender_name} <{sender_email}>",
            "to": [recipient],
            "subject": subject,
            "html": html,
        })

        result_id = getattr(result, "id", None)
        if result_id is None and isinstance(result, dict):
            result_id = result.get("id")

        return jsonify({"ok": True, "id": result_id, "to": recipient})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "10000"))
    app.run(host="0.0.0.0", port=port, debug=False)
